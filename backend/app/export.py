"""Portable reports built from scoped documents, checked facts, and saved answers."""

import csv
import io
import json
from datetime import datetime, timezone
from decimal import Decimal
from html import escape
from threading import Lock

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session
from typing import Literal
from types import SimpleNamespace

from .db import get_db
from .models import Chunk, Document, Fact, QueryRun, ReportBrief

router = APIRouter(prefix="/api/export", tags=["Export"])
_plot_lock = Lock()


class ExportRequest(BaseModel):
    document_ids: list[str] = Field(min_length=1, max_length=50)
    format: Literal["pdf", "xlsx", "md", "csv"] = "pdf"
    include_qa: bool = False
    run_id: str | None = Field(default=None, max_length=32)


def safe_cell(value):
    value = str(value) if value is not None else ""
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def export_facts_csv(facts):
    output = io.StringIO(newline="")
    fields = [
        "company",
        "metric",
        "value",
        "currency",
        "units",
        "period_start",
        "period_end",
        "basis",
        "revision",
        "document_id",
        "page",
        "synthetic",
    ]
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    writer.writerows({key: safe_cell(row.get(key)) for key in fields} for row in facts)
    return output.getvalue().encode("utf-8-sig")


def scoped_facts(db, ids=None):
    query = (
        select(Fact, Document, Chunk)
        .join(Document, Document.id == Fact.document_id)
        .join(Chunk, Chunk.id == Fact.chunk_id)
        .where(Fact.verified.is_(True))
    )
    if ids is not None:
        query = query.where(Document.id.in_(ids))
    return [
        {
            "id": f.id,
            "company": d.company,
            "metric": f.metric,
            "value": f.value,
            "currency": f.currency,
            "units": f.units,
            "period_start": str(f.period_start),
            "period_end": str(f.period_end),
            "basis": f.basis,
            "revision": f.revision,
            "document_id": d.id,
            "page": c.page_index + 1,
            "synthetic": d.synthetic,
        }
        for f, d, c in db.execute(query.order_by(Document.company, Fact.period_end))
    ]


def generate_analysis_report(documents, questions_asked, facts, calculations, format="md"):
    generated = datetime.now(timezone.utc).isoformat()
    lines = [
        "# FilingLens analysis report",
        "",
        f"Generated: {generated}",
        "",
        "## Executive Summary",
        "",
        f"{len(documents)} reports; {len(facts)} checked financial figures. This report contains source excerpts and deterministic calculations. Verify cited pages before relying on the analysis.",
    ]
    if any(d.synthetic for d in documents):
        lines += ["", "**Includes made-up sample data.**"]
    lines += ["", "## Key Metrics", ""]
    lines += [
        f"- {f['company']} — {f['metric']}: {f['value']} {f['currency']} {f['units']}; {f['period_start']} to {f['period_end']}; {f['basis']}, {f['revision']}; source {f['document_id']} page {f['page']}"
        for f in facts
    ] or ["No checked metrics in the selected documents."]
    lines += ["", "## Q&A Results", ""]
    if not questions_asked:
        lines.append("No Q&A history selected.")
    for run in questions_asked:
        lines += [f"### {run.question}", "", run.response.get("summary", "")]
        if run.response.get("abstained"):
            lines.append(run.response["limitation"])
        for claim in run.response.get("claims", []):
            lines.append(claim["text"] + " [" + ", ".join(claim["citation_ids"]) + "]")
    lines += ["", "## Comparative Analysis", ""]
    lines += [
        f"- {c.get('company', '')}: {c['result']} {c['units']}. Formula: {c['formula']}. Inputs: {json.dumps(c.get('normalised_inputs', {}))}"
        for c in calculations
    ] or ["No saved calculations selected."]
    lines += ["", "## Source Citations", ""]
    lines += [
        f"- {d.id}: {d.title} ({d.company}), {d.period_start} to {d.period_end}; {d.page_count} pages."
        for d in documents
    ]
    for run in questions_asked:
        for e in run.response.get("evidence", []):
            lines.append(f"- {e['id']}: {e['title']}, PDF page {e['page_index'] + 1}.")
    if format == "md":
        return "\n".join(lines).encode("utf-8")
    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        summary = wb.active
        summary.title = "Summary"
        summary.append(["FilingLens analysis", "Value"])
        for row in [
            ["Generated", generated],
            ["Reports", len(documents)],
            ["Checked facts", len(facts)],
            ["Includes samples", any(d.synthetic for d in documents)],
        ]:
            summary.append(row)
        sheet = wb.create_sheet("Detailed facts")
        fields = [
            "company",
            "metric",
            "value",
            "currency",
            "units",
            "period_start",
            "period_end",
            "basis",
            "revision",
            "document_id",
            "page",
        ]
        sheet.append(fields)
        for f in facts:
            sheet.append([safe_cell(f[key]) for key in fields])
        calc = wb.create_sheet("Calculations")
        calc.append(
            [
                "Company",
                "Operation",
                "First normalised",
                "Second normalised",
                "Excel formula",
                "Saved result",
                "Units",
                "Formula explanation",
            ]
        )
        for index, c in enumerate(calculations, 2):
            operation = c.get("operation")
            formula = {
                "percentage_change": f"=(D{index}-C{index})/ABS(C{index})*100",
                "difference": f"=D{index}-C{index}",
                "percentage_points": f"=D{index}-C{index}",
                "basis_points": f"=(D{index}-C{index})*100",
            }.get(operation, "")
            inputs = c.get("normalised_inputs", {})
            calc.append(
                [
                    safe_cell(c.get("company")),
                    operation,
                    Decimal(inputs.get("first", "0")),
                    Decimal(inputs.get("second", "0")),
                    formula,
                    c["result"],
                    c["units"],
                    safe_cell(c["formula"]),
                ]
            )
        sources = wb.create_sheet("Sources")
        sources.append(["Document ID", "Report", "Company", "Pages", "Source URL"])
        for d in documents:
            sources.append(
                [
                    safe_cell(d.id),
                    safe_cell(d.title),
                    safe_cell(d.company),
                    d.page_count,
                    safe_cell(d.source_url),
                ]
            )
        sources.append(["Fact page references are in Detailed facts."])
        qa = wb.create_sheet("Q&A")
        qa.append(["Question", "Answer / limitation", "Citations"])
        for run in questions_asked:
            for c in run.response.get("claims", []) or [
                {"text": run.response.get("summary") or run.response["limitation"], "citation_ids": []}
            ]:
                qa.append(
                    [safe_cell(run.question), safe_cell(c["text"]), safe_cell(", ".join(c["citation_ids"]))]
                )
        for ws in wb:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="14634B")
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(
                    55, max(16, max(len(str(c.value or "")) for c in column[:50]) + 2)
                )
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
    if format == "csv":
        return export_facts_csv(facts)
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Image,
        Table,
        TableStyle,
        CondPageBreak,
    )

    output = io.BytesIO()
    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.HexColor("#14634b")
    styles["Normal"].spaceAfter = 6
    styles["Normal"].leading = 14
    story = []
    source_numbers = {d.id: str(i + 1) for i, d in enumerate(documents)}
    table = None
    if facts:
        rows = [["Company", "Metric", "Figure", "Period end", "Basis", "Source"]]
        for f in facts:
            amount = (
                f"{f['value']}%" if f["units"] == "percent" else f"{f['value']} {f['currency']} {f['units']}"
            )
            values = [
                f["company"],
                f["metric"].replace("_", " "),
                amount,
                f["period_end"],
                f["basis"],
                f"[{source_numbers[f['document_id']]}] p. {f['page']}",
            ]
            rows.append([Paragraph(escape(str(value)), styles["Normal"]) for value in values])
        table = Table(rows, colWidths=[62, 87, 94, 82, 85, 60], repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14634b")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f6f3")]),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#dce3eb")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
    in_metrics = False
    for line in lines:
        if line == "## Key Metrics":
            in_metrics = True
        elif line.startswith("## "):
            in_metrics = False
        if in_metrics and line.startswith("- "):
            continue
        if not line:
            story.append(Spacer(1, 8))
        else:
            if line.startswith("## "):
                story.append(CondPageBreak(100))
            style = (
                "Title"
                if line.startswith("# ")
                else "Heading2"
                if line.startswith("## ")
                else "Heading3"
                if line.startswith("### ")
                else "Normal"
            )
            text = line.lstrip("# ").replace("₹", "INR ").replace("—", "-")
            for doc in documents:
                text = text.replace(doc.id, f"[{source_numbers[doc.id]}]")
            text = text.replace(generated, datetime.fromisoformat(generated).strftime("%d %b %Y, %H:%M UTC"))
            text = escape(text)
            if text.startswith("**") and text.endswith("**"):
                text = "<b>" + text[2:-2] + "</b>"
            story.append(Paragraph(text, styles[style]))
            if line == "## Key Metrics" and table is not None:
                story.extend([Spacer(1, 8), table, Spacer(1, 12)])
    revenue = [f for f in facts if f["metric"] == "revenue"]
    if revenue:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        with _plot_lock:
            from collections import defaultdict

            groups = defaultdict(list)
            for fact in revenue:
                groups[
                    (fact["currency"], fact["units"], fact["basis"], fact["revision"], fact["synthetic"])
                ].append(fact)
            fig, axes = plt.subplots(
                min(4, len(groups)), 1, figsize=(7, 3 * min(4, len(groups))), squeeze=False
            )
            for ax, (key, group) in zip(axes[:, 0], list(groups.items())[:4], strict=True):
                for company in sorted({f["company"] for f in group}):
                    series = sorted(
                        [f for f in group if f["company"] == company], key=lambda f: f["period_end"]
                    )
                    ax.plot(
                        [f["period_end"] for f in series],
                        [float(f["value"]) for f in series],
                        marker="o",
                        label=company,
                    )
                ax.set_title(
                    f"Revenue: {key[0]} {key[1]} · {key[2]} · {key[3]}" + (" · synthetic" if key[4] else ""),
                    fontsize=9,
                )
                ax.legend()
                ax.grid(alpha=0.2)
                ax.tick_params(axis="x", labelsize=8)
            fig.tight_layout()
            chart = io.BytesIO()
            fig.savefig(chart, format="png", bbox_inches="tight")
            plt.close(fig)
        chart.seek(0)
        story.append(Image(chart, width=470, height=min(600, 200 * min(4, len(groups)))))
    SimpleDocTemplate(output, title="FilingLens analysis", author="FilingLens").build(story)
    return output.getvalue()


@router.post("/report")
def report(request: ExportRequest, db: Session = Depends(get_db)):
    docs = db.scalars(select(Document).where(Document.id.in_(request.document_ids))).all()
    if len(docs) != len(set(request.document_ids)):
        raise HTTPException(404, "One or more selected documents were not found")
    ids = {d.id for d in docs}
    runs = []
    if request.include_qa or request.run_id:
        statement = (
            select(QueryRun)
            .where(QueryRun.response.is_not(None))
            .order_by(QueryRun.created_at.desc())
            .limit(100)
        )
        if request.run_id:
            statement = statement.where(QueryRun.id == request.run_id)
        for run in db.scalars(statement):
            cited = {e["document_id"] for e in run.response.get("evidence", [])}
            if cited and cited <= ids:
                runs.append(run)
        if request.run_id and not runs:
            raise HTTPException(404, "This answer was not found in the selected reports")
    facts = scoped_facts(db, ids)
    if not request.run_id:
        saved_ids = {r.id for r in runs}
        for brief in db.scalars(
            select(ReportBrief).where(ReportBrief.document_id.in_(ids), ReportBrief.status == "ready")
        ):
            if brief.content.get("run_id") not in saved_ids:
                doc = next(d for d in docs if d.id == brief.document_id)
                runs.append(SimpleNamespace(question=f"Report brief — {doc.title}", response=brief.content))
    calculations = [c for r in runs for c in r.response.get("calculations", [])]
    body = generate_analysis_report(docs, runs, facts, calculations, request.format)
    types = {
        "pdf": "application/pdf",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "md": "text/markdown",
        "csv": "text/csv",
    }
    return Response(
        body,
        media_type=types[request.format],
        headers={"Content-Disposition": f'attachment; filename="filinglens-analysis.{request.format}"'},
    )


@router.get("/facts")
def export_facts(format: Literal["csv", "json"] = "csv", db: Session = Depends(get_db)):
    facts = scoped_facts(db)
    return Response(
        export_facts_csv(facts) if format == "csv" else json.dumps(facts, indent=2).encode(),
        media_type="text/csv" if format == "csv" else "application/json",
        headers={"Content-Disposition": f'attachment; filename="filinglens-facts.{format}"'},
    )
