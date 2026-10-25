import io
from sqlalchemy import select
from app.models import Document


def test_exports_have_sources_and_valid_files(client, corpus):
    ids = list(corpus.scalars(select(Document.id)))
    answer = client.post(
        "/api/questions", json={"question": "Compare revenue growth", "filters": {"synthetic": True}}
    ).json()
    for format in ["md", "pdf", "xlsx", "csv"]:
        response = client.post(
            "/api/export/report", json={"document_ids": ids, "format": format, "include_qa": True}
        )
        assert response.status_code == 200, response.text[:300] if format == "md" else response.status_code
        assert "attachment;" in response.headers["content-disposition"]
        if format == "pdf":
            import pymupdf

            pdf = pymupdf.open(stream=response.content, filetype="pdf")
            text = " ".join(p.get_text() for p in pdf)
            assert "Source Citations" in text and "Key Metrics" in text
        if format == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(response.content))
            assert set(wb.sheetnames) >= {"Summary", "Detailed facts", "Calculations", "Sources"}
            assert wb["Calculations"]["E2"].value.startswith("=")
    assert client.get("/api/export/facts?format=csv").content.startswith(b"\xef\xbb\xbf")
    assert (
        client.post("/api/export/report", json={"document_ids": ["missing"], "format": "md"}).status_code
        == 404
    )
    assert answer["calculations"]
